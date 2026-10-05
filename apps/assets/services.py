import csv
from datetime import datetime
from io import BytesIO, StringIO

from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Asset, AssetCategory, Location

HEADER_FILL = PatternFill("solid", fgColor="0D6EFD")
HEADER_FONT = Font(color="FFFFFF", bold=True)


CSV_COLUMNS = [
    "tag",
    "name",
    "type",
    "status",
    "category",
    "manufacturer",
    "model_name",
    "serial_number",
    "building",
    "floor",
    "room",
    "purchase_date",
    "warranty_until",
    "purchase_price",
    "assigned_to_username",
    "notes",
]


def csv_template():
    buf = StringIO()
    w = csv.writer(buf)
    w.writerow(CSV_COLUMNS)
    w.writerow(
        [
            "",
            "Laptop Dell",
            "LAPTOP",
            "IN_USE",
            "Sprzęt",
            "Dell",
            "Latitude 5540",
            "SN123456",
            "Budynek A",
            "1",
            "101",
            "2024-01-15",
            "2027-01-15",
            "4599.00",
            "admin",
            "Przykładowy wiersz — usuń przed importem",
        ]
    )
    return buf.getvalue()


def _parse_date(v):
    v = (v or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    return None


@transaction.atomic
def import_assets_from_csv(file_obj, user):
    from apps.accounts.models import User

    data = file_obj.read()
    if isinstance(data, bytes):
        data = data.decode("utf-8-sig")

    reader = csv.DictReader(StringIO(data))
    created = updated = 0
    errors = []

    if reader.fieldnames:
        missing = set(["tag", "name"]) - set(reader.fieldnames)
        if missing:
            errors.append(f"Brak wymaganych kolumn: {', '.join(missing)}")
            return created, updated, errors

    for i, row in enumerate(reader, start=2):
        try:
            tag = (row.get("tag") or "").strip()
            name = (row.get("name") or "").strip()
            if not name:
                errors.append(f"Wiersz {i}: brak nazwy — pominięty.")
                continue

            defaults = {
                "name": name,
                "type": (row.get("type") or "OTHER").strip() or "OTHER",
                "status": (row.get("status") or "IN_STOCK").strip() or "IN_STOCK",
                "manufacturer": (row.get("manufacturer") or "").strip(),
                "model_name": (row.get("model_name") or "").strip(),
                "serial_number": (row.get("serial_number") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
                "purchase_date": _parse_date(row.get("purchase_date")),
                "warranty_until": _parse_date(row.get("warranty_until")),
            }

            price = (row.get("purchase_price") or "").strip().replace(",", ".")
            defaults["purchase_price"] = price or None

            cat_name = (row.get("category") or "").strip()
            if cat_name:
                cat, _ = AssetCategory.objects.get_or_create(name=cat_name)
                defaults["category"] = cat

            building = (row.get("building") or "").strip()
            if building:
                loc, _ = Location.objects.get_or_create(
                    building=building,
                    floor=(row.get("floor") or "").strip(),
                    room=(row.get("room") or "").strip(),
                )
                defaults["location"] = loc

            username = (row.get("assigned_to_username") or "").strip()
            if username:
                u = User.objects.filter(username=username).first()
                if u:
                    defaults["assigned_to"] = u
                else:
                    errors.append(f"Wiersz {i}: nie znaleziono użytkownika {username}.")

            if tag:
                _, is_new = Asset.objects.update_or_create(tag=tag, defaults=defaults)
            else:
                Asset.objects.create(**defaults)
                is_new = True

            if is_new:
                created += 1
            else:
                updated += 1

        except Exception as e:
            errors.append(f"Wiersz {i}: {e}")

    return created, updated, errors


def export_assets_to_xlsx(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Majątek"

    headers = [
        "Tag",
        "Nazwa",
        "Typ",
        "Status",
        "Kategoria",
        "Producent",
        "Model",
        "SN",
        "Budynek",
        "Piętro",
        "Pokój",
        "Przypisany do",
        "Data zakupu",
        "Gwarancja do",
        "Cena",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for a in queryset.select_related("category", "location", "assigned_to"):
        ws.append(
            [
                a.tag,
                a.name,
                a.get_type_display(),
                a.get_status_display(),
                a.category.name if a.category else "",
                a.manufacturer,
                a.model_name,
                a.serial_number,
                a.location.building if a.location else "",
                a.location.floor if a.location else "",
                a.location.room if a.location else "",
                a.assigned_to.username if a.assigned_to else "",
                a.purchase_date.isoformat() if a.purchase_date else "",
                a.warranty_until.isoformat() if a.warranty_until else "",
                str(a.purchase_price) if a.purchase_price else "",
            ]
        )

    widths = [16, 32, 14, 14, 20, 16, 20, 18, 14, 8, 8, 16, 14, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
