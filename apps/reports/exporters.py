from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="0D6EFD")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def export_report_to_xlsx(report):
    wb = Workbook()

    # Arkusz 1 — podsumowanie
    ws = wb.active
    ws.title = "Podsumowanie"
    ws.append(["Metryka", "Wartość"])
    _style_header(ws)
    ws.append(["Zakres od", report["date_from"].isoformat()])
    ws.append(["Zakres do", report["date_to"].isoformat()])
    ws.append(["Liczba zgłoszeń", report["total"]])
    ws.append(["Rozwiązane", report["resolved"]])
    ws.append(["Średni czas rozwiązania (h)", report["avg_resolve_hours"] or "—"])
    ws.append(["SLA przekroczone", report["sla_breached"]])
    ws.append(["SLA % (dotrzymane)", report["sla_percent"]])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    # Arkusz 2 — per kategoria
    ws2 = wb.create_sheet("Kategorie")
    ws2.append(["Kategoria", "Liczba zgłoszeń"])
    _style_header(ws2)
    for row in report["by_category"]:
        ws2.append([row["category__name"] or "Bez kategorii", row["total"]])
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 18

    # Arkusz 3 — per priorytet
    ws3 = wb.create_sheet("Priorytety")
    ws3.append(["Priorytet", "Liczba zgłoszeń"])
    _style_header(ws3)
    for row in report["by_priority"]:
        ws3.append([row["priority"], row["total"]])
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 18

    # Arkusz 4 — per technik
    ws4 = wb.create_sheet("Technicy")
    ws4.append(["Technik", "Liczba zgłoszeń"])
    _style_header(ws4)
    for row in report["by_technician"]:
        name = (
            " ".join(filter(None, [row["assigned_to__first_name"], row["assigned_to__last_name"]]))
            or row["assigned_to__username"]
            or "— nieprzypisane —"
        )
        ws4.append([name, row["total"]])
    ws4.column_dimensions["A"].width = 32
    ws4.column_dimensions["B"].width = 18

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
